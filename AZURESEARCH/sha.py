from fastapi import FastAPI, HTTPException
from pydantic import BaseModel, Field
from typing import Optional, List, Dict
import requests
import pdfplumber
from docx import Document
import io
import os
import csv
import openpyxl
from azure.search.documents import SearchClient
from azure.search.documents.indexes import SearchIndexClient
from azure.search.documents.indexes.models import (SearchIndex, SimpleField, SearchField, SearchableField, 
                                                   HnswAlgorithmConfiguration, VectorSearch, 
                                                   VectorSearchAlgorithmConfiguration, VectorSearchProfile, SearchFieldDataType)
from azure.core.credentials import AzureKeyCredential
from openai import AzureOpenAI
from azure.identity import DefaultAzureCredential, get_bearer_token_provider
import logging

# Initialize FastAPI
app = FastAPI()

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Azure and OpenAI configurations
search_service_endpoint = os.getenv("SEARCH_SERVICE_ENDPOINT", "https://your-search-service.search.windows.net")
admin_key = os.getenv("SEARCH_SERVICE_ADMIN_KEY", "your-search-service-admin-key")
index_name = os.getenv("SEARCH_INDEX_NAME", "sharepoint")
group_id = os.getenv("GROUP_ID", "your-group-id")
azure_openai_endpoint = os.getenv("AZURE_OPENAI_ENDPOINT", "https://your-openai-service.openai.azure.com/")
azure_openai_key = os.getenv("AZURE_OPENAI_KEY", "your-openai-api-key")
azure_openai_embedding_deployment = os.getenv("AZURE_OPENAI_EMBEDDING_DEPLOYMENT", "text-embedding-ada-002")
embedding_model_name = "text-embedding-ada-002"
azure_openai_version = "2023-05-15"

# Define request models
class SharePointItem(BaseModel):
    ACCESS_TOKEN: Optional[str] = Field(None, description="Access token for SharePoint API")
    CLIENT_ID: Optional[str] = Field(None, description="Client ID for SharePoint API")
    CLIENT_SECRET: Optional[str] = Field(None, description="Client secret for SharePoint API")
    TENANT_ID: Optional[str] = Field(None, description="Tenant ID for SharePoint API")
    configuration: str = Field(..., description="Configuration type: 'list', 'drive', or 'pages'")
    SHAREPOINT_DRIVE_ID: Optional[str] = Field(None, description="ID of the SharePoint drive (optional)")
    list_id: Optional[str] = Field(None, description="ID of the SharePoint list (optional)")
    sharepoint_site_id: Optional[str] = Field(None, description="ID of the SharePoint site (optional)")
    folder_id: Optional[str] = Field(None, description="ID of the SharePoint folder (optional)")

# Helper functions
def sanitize_field_name(field_name):
    return field_name[1:] if field_name.startswith('_') else field_name

def sanitize_field_names(fields):
    return [sanitize_field_name(field) for field in fields]

def get_access_token(client_id, client_secret, tenant_id):
    token_url = f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token"
    data = {
        "client_id": client_id,
        "client_secret": client_secret,
        "scope": "https://graph.microsoft.com/.default",
        "grant_type": "client_credentials"
    }

    response = requests.post(token_url, data=data)
    if response.status_code == 200:
        return response.json().get("access_token")
    else:
        logger.error(f"Failed to retrieve access token: {response.text}")
        return None

def get_sharepoint_pages(site_id, access_token):
    graph_endpoint = f"https://graph.microsoft.com/v1.0/sites/{site_id}/pages"
    headers = {"Authorization": f"Bearer {access_token}"}
    response = requests.get(graph_endpoint, headers=headers)
    if response.status_code == 200:
        return response.json().get("value", [])
    else:
        logger.error(f"Failed to retrieve SharePoint pages data: {response.text}")
        return []

def create_or_update_index(fields):
    client = SearchIndexClient(endpoint=search_service_endpoint, credential=AzureKeyCredential(admin_key))
    field_schemas = [SimpleField(name=sanitize_field_name(field), type="Edm.String", filterable=True) for field in fields]
    field_schemas.append(SimpleField(name="id", type=SearchFieldDataType.String, key=True))
    field_schemas.append(SimpleField(name="accessList", type="Collection(Edm.String)"))
    field_schemas.append(SimpleField(name="permissions", type="Collection(Edm.String)"))
    field_schemas.append(SearchField(name="contentEmbeddings", type="Collection(Edm.Single)", vector_search_dimensions=1536, vector_search_profile_name="myHnswProfile"))

    vector_search = VectorSearch(
        algorithms=[HnswAlgorithmConfiguration(name="myHnsw")],
        profiles=[VectorSearchProfile(name="myHnswProfile", algorithm_configuration_name="myHnsw")]
    )

    index = SearchIndex(name=index_name, fields=field_schemas, vector_search=vector_search)
    try:
        existing_index = client.get_index(index_name)
        logger.info(f"Index '{index_name}' already exists. Deleting the existing index...")
        client.delete_index(index_name)
        logger.info(f"Index '{index_name}' has been deleted.")
    except Exception as e:
        logger.info(f"Index '{index_name}' not found. Creating a new one.")

    try:
        logger.info(f"Creating index '{index_name}'...")
        client.create_index(index)
        logger.info(f"Index '{index_name}' has been created successfully.")
    except Exception as e:
        logger.error(f"Error creating index '{index_name}': {e}")

def download_sharepoint_file(file_id, sharepoint_drive_id, access_token):
    content_url = f"https://graph.microsoft.com/v1.0/drives/{sharepoint_drive_id}/items/{file_id}/content"
    headers = {"Authorization": "Bearer " + access_token}
    response = requests.get(content_url, headers=headers)
    if response.status_code == 200:
        return response.content
    else:
        raise HTTPException(status_code=response.status_code, detail=f"Failed to download file with ID '{file_id}'")

def get_sharepoint_items(sharepoint_drive_id, access_token, folder_id=None):
    url = f"https://graph.microsoft.com/v1.0/drives/{sharepoint_drive_id}/items/{folder_id}/children" if folder_id else f"https://graph.microsoft.com/v1.0/drives/{sharepoint_drive_id}/root/children"
    headers = {"Authorization": "Bearer " + access_token, "Accept": "application/json"}
    response = requests.get(url, headers=headers)
    return response.json().get('value', [])

def get_permissions(sharepoint_drive_id, file_id, access_token):
    permissions_url = f"https://graph.microsoft.com/v1.0/drives/{sharepoint_drive_id}/items/{file_id}/permissions"
    headers = {"Authorization": "Bearer " + access_token, "Accept": "application/json"}
    response = requests.get(permissions_url, headers=headers)
    permissions = []
    if response.status_code == 200:
        permissions_data = response.json().get("value", [])
        for permission_data in permissions_data:
            granted_to = permission_data.get("grantedTo", {})
            if "user" in granted_to:
                user_or_group = granted_to["user"].get("displayName", "N/A")
                if user_or_group != "N/A":
                    permissions.append(user_or_group)
            elif "group" in granted_to:
                user_or_group = granted_to["group"].get("displayName", "N/A")
                if user_or_group != "N/A":
                    permissions.append(user_or_group)
    return permissions

def get_group_members(group_id, access_token):
    url = f"https://graph.microsoft.com/v1.0/groups/{group_id}/transitiveMembers/"
    headers = {"Authorization": "Bearer " + access_token, "Accept": "application/json"}
    response = requests.get(url, headers=headers)
    group_members = []
    if response.status_code == 200:
        members = response.json().get("value", [])
        for member in members:
            principal_name = member.get("userPrincipalName", "N/A")
            if principal_name != "N/A":
                group_members.append(principal_name)
    else:
        logger.error(f"Failed to retrieve the group members: {response.text}")
    return group_members

def extract_text_from_pdf(file_content):
    with pdfplumber.open(io.BytesIO(file_content)) as pdf:
        text = "".join(page.extract_text() for page in pdf.pages)
    return text

def extract_text_from_word(file_content):
    doc = Document(io.BytesIO(file_content))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)

def extract_text_from_excel(file_content):
    workbook = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    text = ""
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text += " ".join(str(cell) for cell in row if cell is not None) + "\n"
    return text

def extract_text_from_csv(file_content):
    csv_content = io.StringIO(file_content.decode("utf-8"))
    reader = csv.reader(csv_content)
    return "\n".join(" ".join(row) for row in reader)

def extract_text_from_file(file_content, file_extension):
    if file_extension == ".pdf":
        return extract_text_from_pdf(file_content)
    elif file_extension == ".docx":
        return extract_text_from_word(file_content)
    elif file_extension == ".xlsx":
        return extract_text_from_excel(file_content)
    elif file_extension == ".csv":
        return extract_text_from_csv(file_content)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file extension: {file_extension}")

def generate_embeddings(texts, azure_openai_endpoint, azure_openai_key, deployment_id):
    azure_openai = AzureOpenAI(endpoint=azure_openai_endpoint, api_key=azure_openai_key)
    embeddings = []
    for text in texts:
        response = azure_openai.Embeddings.create(input=[text], deployment_id=deployment_id)
        embeddings.append(response["data"][0]["embedding"])
    return embeddings

def index_documents(documents, search_service_endpoint, admin_key, index_name):
    client = SearchClient(endpoint=search_service_endpoint, index_name=index_name, credential=AzureKeyCredential(admin_key))
    batch = []
    for document in documents:
        try:
            client.upload_documents(documents=[document])
            logger.info(f"Document indexed successfully: {document['id']}")
        except Exception as e:
            logger.error(f"Error indexing document {document['id']}: {e}")
            batch.append(document)
    if batch:
        client.upload_documents(documents=batch)
def get_sharepoint_list_items(list_id, access_token, sharepoint_site_id):
    try:
        schema_url = f"https://graph.microsoft.com/v1.0/sites/{sharepoint_site_id}/lists/{list_id}/columns"
        headers = {
            "Authorization": "Bearer " + access_token,
            "Accept": "application/json"
        }
        response = requests.get(schema_url, headers=headers)
        if response.status_code == 200:
            schema_data = response.json()
            fields_to_expand_and_index = ([column["name"] for column in schema_data["value"]])

            select_fields = ','.join(fields_to_expand_and_index)
            url = f"https://graph.microsoft.com/v1.0/sites/{sharepoint_site_id}/lists/{list_id}/items?expand=fields(select={select_fields})"
            response = requests.get(url, headers=headers)
            if response.status_code == 200:
                data = response.json()
                if 'value' in data:
                    return data['value'], fields_to_expand_and_index
                else:
                    print("No items found in the SharePoint list.")
                    return [], []
            else:
                print(f"Failed to retrieve data from SharePoint list. Status code: {response.status_code}")
                return [], []
        else:
            print(f"Failed to retrieve schema from SharePoint list. Status code: {response.status_code}")
            return [], []
    except Exception as e:
        print(f"Error fetching data from SharePoint list: {str(e)}")
        return [], []

@app.post("/index_sharepoint_items")
def index_sharepoint_items(item: SharePointItem):
    try:
        access_token = item.ACCESS_TOKEN or get_access_token(item.CLIENT_ID, item.CLIENT_SECRET, item.TENANT_ID)
        group_members = get_group_members(group_id, access_token)
        if item.configuration == "list":
            list_items = get_sharepoint_list_items(item.sharepoint_site_id, item.list_id, access_token)
            documents = []
            for list_item in list_items:
                content = list_item.get("fields", {}).get("Content", "")
                embedding = generate_embeddings([content], azure_openai_endpoint, azure_openai_key, azure_openai_embedding_deployment)
                document = {
                    "id": list_item["id"],
                    "content": content,
                    "accessList": group_members,
                    "permissions": [],
                    "contentEmbeddings": embedding[0]
                }
                documents.append(document)
            index_documents(documents, search_service_endpoint, admin_key, index_name)
            return {"status": "success", "message": "SharePoint list items indexed successfully"}
        elif item.configuration == "drive":
            items = get_sharepoint_items(item.SHAREPOINT_DRIVE_ID, access_token, item.folder_id)
            documents = []
            for drive_item in items:
                file_id = drive_item["id"]
                file_name = drive_item["name"]
                file_extension = os.path.splitext(file_name)[1].lower()
                file_content = download_sharepoint_file(file_id, item.SHAREPOINT_DRIVE_ID, access_token)
                extracted_text = extract_text_from_file(file_content, file_extension)
                embedding = generate_embeddings([extracted_text], azure_openai_endpoint, azure_openai_key, azure_openai_embedding_deployment)
                permissions = get_permissions(item.SHAREPOINT_DRIVE_ID, file_id, access_token)
                document = {
                    "id": file_id,
                    "content": extracted_text,
                    "accessList": group_members,
                    "permissions": permissions,
                    "contentEmbeddings": embedding[0]
                }
                documents.append(document)
            index_documents(documents, search_service_endpoint, admin_key, index_name)
            return {"status": "success", "message": "SharePoint drive items indexed successfully"}
        elif item.configuration == "pages":
            pages = get_sharepoint_pages(item.sharepoint_site_id, access_token)
            documents = []
            for page in pages:
                content = page.get("content", "")
                embedding = generate_embeddings([content], azure_openai_endpoint, azure_openai_key, azure_openai_embedding_deployment)
                document = {
                    "id": page["id"],
                    "content": content,
                    "accessList": group_members,
                    "permissions": [],
                    "contentEmbeddings": embedding[0]
                }
                documents.append(document)
            index_documents(documents, search_service_endpoint, admin_key, index_name)
            return {"status": "success", "message": "SharePoint pages indexed successfully"}
        else:
            raise HTTPException(status_code=400, detail="Invalid configuration type")
    except Exception as e:
        logger.error(f"Error indexing SharePoint items: {e}")
        raise HTTPException(status_code=500, detail="An error occurred while indexing SharePoint items")
