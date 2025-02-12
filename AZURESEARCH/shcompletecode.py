from fastapi import FastAPI, HTTPException
from pydantic import BaseModel, Field
from typing import Optional, Dict
import requests
import uvicorn
import pdfplumber
from docx import Document
import io
import os
import csv
import openpyxl
import base64
from azure.search.documents import SearchClient
from azure.search.documents.indexes import SearchIndexClient
from azure.search.documents.indexes.models import SearchIndex, SimpleField, SearchableField, SearchFieldDataType
from azure.core.credentials import AzureKeyCredential
from azure.identity import DefaultAzureCredential
app = FastAPI()

# Define request models
class SharePointItem(BaseModel):
    ACCESS_TOKEN: Optional[str] = Field(..., description="Access token for SharePoint API")
    CLIENT_ID: Optional[str] = Field(None, description="Client ID for SharePoint API")
    CLIENT_SECRET: Optional[str] = Field(None, description="Client secret for SharePoint API")
    TENANT_ID: Optional[str] = Field(None, description="Tenant ID for SharePoint API")
    configuration: str = Field(..., description="Configuration type: 'list', 'drive', or 'pages'")
    SHAREPOINT_DRIVE_ID: Optional[str] = Field(None, description="ID of the SharePoint drive (optional)")
    list_id: Optional[str] = Field(None, description="ID of the SharePoint list (optional)")
    sharepoint_site_id: Optional[str] = Field(None, description="ID of the SharePoint site (optional)")
    folder_id: Optional[str] = Field(None, description="ID of the SharePoint folder (optional)")

def get_access_token(client_id, client_secret, tenant_id):
    token_url = f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token"
    data = {
        "client_id": client_id,
        "client_secret": client_secret,
        "scope": "https://graph.microsoft.com/.default",
        "grant_type": "client_credentials"
    }

    response = requests.post(token_url, data=data)
    if response.status_code == 200:
        return response.json().get("access_token")
    else:
        print(f"Failed to retrieve access token: {response.text}")
        return None
def get_sharepoint_list_data(list_id, access_token, sharepoint_site_id):
    try:
        schema_url = f"https://graph.microsoft.com/v1.0/sites/{sharepoint_site_id}/lists/{list_id}/columns"
        headers = {
            "Authorization": "Bearer " + access_token,
            "Accept": "application/json"
        }
        response = requests.get(schema_url, headers=headers)
        if response.status_code == 200:
            schema_data = response.json()
            fields_to_expand_and_index = ([column["name"] for column in schema_data["value"]])

            select_fields = ','.join(fields_to_expand_and_index)
            url = f"https://graph.microsoft.com/v1.0/sites/{sharepoint_site_id}/lists/{list_id}/items?expand=fields(select={select_fields})"
            response = requests.get(url, headers=headers)
            if response.status_code == 200:
                data = response.json()
                if 'value' in data:
                    return data['value'], fields_to_expand_and_index
                else:
                    print("No items found in the SharePoint list.")
                    return [], []
            else:
                print(f"Failed to retrieve data from SharePoint list. Status code: {response.status_code}")
                return [], []
        else:
            print(f"Failed to retrieve schema from SharePoint list. Status code: {response.status_code}")
            return [], []
    except Exception as e:
        print(f"Error fetching data from SharePoint list: {str(e)}")
        return [], []

def get_all_sharepoint_lists(access_token, sharepoint_site_id):
    try:
        lists_url = f"https://graph.microsoft.com/v1.0/sites/{sharepoint_site_id}/lists"
        print("retriveing list url {list_url}")
        headers = {
            "Authorization": "Bearer " + access_token,
            "Accept": "application/json"
        }
        response = requests.get(lists_url, headers=headers)
        if response.status_code == 200:
            lists_data = response.json().get("value", [])
            all_lists_data = []
            for list_item in lists_data:
                list_id = list_item.get("id")
                list_data, fields = get_sharepoint_list_data(list_id, access_token, sharepoint_site_id)
                all_lists_data.append({"list_id": list_id, "list_data": list_data, "fields": fields})
            return all_lists_data
        else:
            print(f"Failed to retrieve lists from SharePoint site. Status code: {response.status_code}")
            return []
    except Exception as e:
        print(f"Error fetching lists from SharePoint site: {str(e)}")
        return []

def get_sharepoint_items(access_token, folder_id, SHAREPOINT_DRIVE_ID):
    try:
        if folder_id:
            url = f"https://graph.microsoft.com/v1.0/drives/{SHAREPOINT_DRIVE_ID}/items/{folder_id}/children"
        else:
            url = f"https://graph.microsoft.com/v1.0/drives/{SHAREPOINT_DRIVE_ID}/root/children"

        headers = {
            "Authorization": "Bearer " + access_token,
            "Accept": "application/json"
        }
        response = requests.get(url, headers=headers)
        if response.status_code == 200:
            data = response.json()['value']
            return data
        else:
            raise HTTPException(status_code=response.status_code, detail="Failed to retrieve items from SharePoint drive.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error fetching items from SharePoint drive: {str(e)}")

def download_sharepoint_file(file_id, SHAREPOINT_DRIVE_ID, access_token):
    try:
        content_url = f"https://graph.microsoft.com/v1.0/drives/{SHAREPOINT_DRIVE_ID}/items/{file_id}/content"
        headers = {
            "Authorization": "Bearer " + access_token,
        }
        response = requests.get(content_url, headers=headers)
        if response.status_code == 200:
            return response.content
        else:
            raise HTTPException(status_code=response.status_code, detail=f"Failed to download file with ID '{file_id}'")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error downloading file from SharePoint drive: {str(e)}")

def extract_text_from_pdf(file_content):
    with pdfplumber.open(io.BytesIO(file_content)) as pdf:
        return "\n".join(page.extract_text() for page in pdf.pages)

def extract_text_from_word(file_content):
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(file_content)).paragraphs)

def extract_text_from_excel(file_content):
    file_str = ""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content))
        for sheet in wb:
            file_str += " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    except Exception as e:
        file_str = f"Error extracting text from Excel file: {str(e)}"
    return file_str

def extract_text_from_text(file_content):
    try:
        return file_content.decode('utf-8')
    except UnicodeDecodeError:
        return str(file_content, errors='ignore')

def extract_text_from_csv(file_content):
    try:
        return " ".join(" ".join(row) for row in csv.reader(io.StringIO(file_content.decode('utf-8'))))
    except Exception as e:
        return f"Error extracting text from CSV file: {str(e)}"

def extract_text_from_other(file_content):
    return "Unsupported file type"

def extract_text(file_name, file_content):
    ext = os.path.splitext(file_name)[-1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_content)
    elif ext in [".doc", ".docx"]:
        return extract_text_from_word(file_content)
    elif ext in [".xls", ".xlsx"]:
        return extract_text_from_excel(file_content)
    elif ext == ".txt":
        return extract_text_from_text(file_content)
    elif ext == ".csv":
        return extract_text_from_csv(file_content)
    else:
        return extract_text_from_other(file_content)
def get_sharepoint_pages(sharepoint_site_id, access_token):
    graph_endpoint = f"https://graph.microsoft.com/v1.0/sites/{sharepoint_site_id}/pages"
    headers = {
        "Authorization": f"Bearer {access_token}"
    }
    response = requests.get(graph_endpoint, headers=headers)
    if response.status_code == 200:
        return response.json().get("value", [])
    else:
        print(f"Failed to retrieve SharePoint pages data: {response.text}")
        return []


@app.post("/sharepoint/get-details/")
def get_sharepoint_details(item: SharePointItem):
    if not item.ACCESS_TOKEN:
        if not (item.CLIENT_ID and item.CLIENT_SECRET and item.TENANT_ID):
            raise HTTPException(status_code=400, detail="Access token or client credentials must be provided.")
        item.ACCESS_TOKEN = get_access_token(item.CLIENT_ID, item.CLIENT_SECRET, item.TENANT_ID)

    if item.configuration == "list":
        if item.list_id and item.sharepoint_site_id:
            list_data, fields = get_sharepoint_list_data(item.list_id, item.ACCESS_TOKEN, item.sharepoint_site_id)
            if not list_data:
                raise HTTPException(status_code=404, detail="No items found in the SharePoint list.")
           
            return {"list_data": list_data, "fields": fields}
        else:
            all_lists_data = get_all_sharepoint_lists(item.ACCESS_TOKEN, item.sharepoint_site_id)
            if not all_lists_data:
                raise HTTPException(status_code=404, detail="No lists found in the SharePoint site.")
            for list_data_info in all_lists_data:
                list_id = list_data_info["list_id"]
                list_data = list_data_info["list_data"]
                fields = list_data_info["fields"]
                
                
                
            
            return {"all_lists_data": all_lists_data}
            
    elif item.configuration == "drive":
        if item.SHAREPOINT_DRIVE_ID:
            if item.folder_id:
                drive_items = get_sharepoint_items(item.ACCESS_TOKEN, item.folder_id, item.SHAREPOINT_DRIVE_ID)
                file_contents = []
                for drive_item in drive_items:
                    if drive_item.get('file'):
                        file_id = drive_item['id']
                        file_name = drive_item['name']
                        file_content = download_sharepoint_file(file_id, item.SHAREPOINT_DRIVE_ID, item.ACCESS_TOKEN)
                        text_content = extract_text(file_name, file_content)
                        file_contents.append({"fileName": file_name, "content": text_content})
                return {"file_contents": file_contents}
            else:
                drive_items = get_sharepoint_items(item.ACCESS_TOKEN, None, item.SHAREPOINT_DRIVE_ID)
                file_contents = []
                for drive_item in drive_items:
                    if drive_item.get('file'):
                        file_id = drive_item['id']
                        file_name = drive_item['name']
                        file_content = download_sharepoint_file(file_id, item.SHAREPOINT_DRIVE_ID, item.ACCESS_TOKEN)
                        text_content = extract_text(file_name, file_content)
                        file_contents.append({"fileName": file_name, "content": text_content})
                return {"file_contents": file_contents}
        else:
            raise HTTPException(status_code=400, detail="Drive ID is required for 'drive' configuration.")
    elif item.configuration == "pages":
        if item.sharepoint_site_id:
            pages = get_sharepoint_pages(item.sharepoint_site_id, item.ACCESS_TOKEN)
            print("pages de",pages)
            file_contents = []
            for page in pages:
                file_contents.append({
                    "id": str(page["id"]),
                    "title": page.get("title", ""),
                    "weburl": page.get("webUrl", "")
                })
                return {"file_contents": file_contents}
        else:
            raise HTTPException(status_code=400, detail="SharePoint site ID is required for 'pages' configuration.")
    else:
        raise HTTPException(status_code=400, detail="Invalid configuration type. Must be 'list', 'drive', or 'pages'.")

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)
